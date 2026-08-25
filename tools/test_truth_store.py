"""Tests for the accumulating truth store.

The operator's complaint this exists to answer: "the info I provide on shots should be saved
as truths and built upon so I don't have to keep reviewing the same info." So the failures
that matter are losing a fact, or letting a weaker source overwrite a stronger one.
"""
from __future__ import annotations

from tools import truth_store as ts


def test_a_better_source_supersedes_and_the_change_is_kept():
    """Import order used to decide the winner, so a 2026-08-17 spreadsheet overruled a review
    done today. Authority decides, and what was replaced stays visible."""
    doc = ts.empty("v.mp4")
    ts.add_shot(doc, 10.0, type_="drive", source="old.csv", kind="labels_csv")
    v = ts.add_shot(doc, 10.0, type_="drop", source="review", kind="review")
    assert v == "updated"
    s = doc["shots"][0]
    assert s["type"] == "drop"
    assert s["superseded"][0] == {"field": "type", "was": "drive", "now": "drop",
                                  "by": "review"}


def test_a_weaker_source_cannot_overwrite():
    doc = ts.empty("v.mp4")
    ts.add_shot(doc, 10.0, type_="drop", source="review", kind="review")
    ts.add_shot(doc, 10.0, type_="drive", source="old.csv", kind="labels_csv")
    assert doc["shots"][0]["type"] == "drop"
    assert "superseded" not in doc["shots"][0]


def test_equal_authority_disagreement_is_a_conflict_not_a_silent_pick():
    doc = ts.empty("v.mp4")
    ts.add_shot(doc, 10.0, type_="drop", source="a", kind="review")
    v = ts.add_shot(doc, 10.0, type_="lob", source="b", kind="review")
    assert v == "CONFLICT"
    assert doc["shots"][0]["type"] == "drop"          # first kept, both recorded
    assert doc["shots"][0]["conflicts"][0]["rejected"] == "lob"


def test_one_stored_shot_cannot_absorb_several_distinct_ones():
    """A +/-1 s window is right for hand-typed times and wide enough to span three shots of a
    kitchen exchange. Without one-to-one claiming, one row swallowed them all and then
    'conflicted' with every one -- 117 false conflicts on the first import."""
    doc = ts.empty("v.mp4")
    claimed: set = set()
    for t, ty in ((10.0, "dink"), (10.4, "dink"), (10.8, "drive")):
        ts.add_shot(doc, t, type_=ty, source="review", kind="review", claimed=claimed)
    assert len(doc["shots"]) == 3
    assert not any(s.get("conflicts") for s in doc["shots"])


def test_hitter_naming_is_normalised():
    """The operator writes 'opponent'; track_roles says opp_a. Comparing those raw invents a
    disagreement out of a naming difference."""
    assert ts.norm_hitter("opp_a") == ts.norm_hitter("opponent") == "opponent"
    assert ts.norm_hitter("Partner") == "partner"
    assert ts.norm_hitter("") is None
    doc = ts.empty("v.mp4")
    ts.add_shot(doc, 5.0, hitter="opp_b", source="a", kind="labels_csv")
    v = ts.add_shot(doc, 5.0, hitter="opponent", source="b", kind="review")
    assert v == "agreed"


def test_enrichment_adds_without_conflicting():
    doc = ts.empty("v.mp4")
    ts.add_shot(doc, 5.0, type_="dink", source="a", kind="labels_csv")
    v = ts.add_shot(doc, 5.0, hitter="user", side="near", source="b", kind="labels_csv")
    assert v == "enriched"
    s = doc["shots"][0]
    assert s["type"] == "dink" and s["hitter"] == "user" and s["side"] == "near"


def test_store_is_keyed_by_video_not_by_folder():
    """Eight analysed folders exist for one video. Truth belongs to the video, or it is
    re-collected every time a clip is re-analysed."""
    a = ts.store_path("PB 5 minute outdoor.mp4")
    b = ts.store_path("PB 5 minute outdoor.mp4")
    assert a == b and a.suffix == ".json"
    assert ts.store_path("PB 3 min indoor 1 court B.mp4") != a


def test_an_untouched_review_sheet_is_not_imported_as_truth(tmp_path):
    """A blank row means "the operator agreed" -- but only in a sheet they worked through.

    Importing a freshly built sheet would file OUR OWN detections as operator truth at the
    highest authority, and every accuracy figure for that clip would be us scoring ourselves.
    """
    import tools.shot_review_sheet as srs
    from tools.test_shot_review_sheet import _clip
    clip = _clip(tmp_path)
    srs.build(clip, clip / "_labeling" / srs.OUT_NAME)
    doc = ts.empty("v.mp4")
    assert ts.import_review_xlsx(doc, clip) == {}
    assert doc["shots"] == []

    # ...and one mark anywhere in it is enough to make the whole sheet count.
    from openpyxl import load_workbook
    wb = load_workbook(clip / "_labeling" / srs.OUT_NAME)
    ws = wb.active
    hdr = next(r for r in range(1, 30) if ws.cell(row=r, column=1).value == "#")
    col = {str(ws.cell(row=hdr, column=i).value or "").strip(): i
           for i in range(1, ws.max_column + 1)}
    ws.cell(row=hdr + 2, column=col["CORRECT_TYPE"], value="drop")
    wb.save(clip / "_labeling" / srs.OUT_NAME)
    doc2 = ts.empty("v.mp4")
    ts.import_review_xlsx(doc2, clip)
    assert doc2["shots"]


def test_importing_the_same_review_twice_changes_nothing(tmp_path):
    """The store accumulates the operator's answers, so it must be a function of its sources.

    It was not: re-running --import-all appended clones of shots whose neighbours sat inside
    the +/-1s match window, because each of the two rows claimed the other's entry and the
    loser appended a fresh one. Eight duplicate pairs had built up that way, and a duplicated
    shot inflates the total and every rate computed from it.
    """
    import json
    import tools.shot_review_sheet as srs
    from tools.test_shot_review_sheet import _clip
    from openpyxl import load_workbook
    clip = _clip(tmp_path)
    sheet = clip / "_labeling" / srs.OUT_NAME
    srs.build(clip, sheet)
    wb = load_workbook(sheet)
    ws = wb.active
    hdr = next(r for r in range(1, 30) if ws.cell(row=r, column=1).value == "#")
    col = {str(ws.cell(row=hdr, column=i).value or "").strip(): i
           for i in range(1, ws.max_column + 1)}
    ws.cell(row=hdr + 2, column=col["CORRECT_TYPE"], value="drop")
    # two shots inside one match window: the case that produced the clones
    blank = hdr + 3
    ws.cell(row=blank, column=col["time"], value="0:01.40")
    ws.cell(row=blank, column=col["CORRECT_TYPE"], value="dink")
    wb.save(sheet)

    doc = ts.empty("v.mp4")
    ts.import_review_xlsx(doc, clip)
    ts.fold_shadowed_legacy(doc)
    once = json.dumps(doc, sort_keys=True)
    ts.import_review_xlsx(doc, clip)
    ts.fold_shadowed_legacy(doc)
    assert json.dumps(doc, sort_keys=True) == once
