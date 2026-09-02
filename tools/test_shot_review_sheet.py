"""Tests for the shot-review workbook round-trip.

The operator's time is the scarce input, so the failure that matters is a correction being
written and then not making it into the truth set -- silently, since a missing label looks
exactly like a shot nobody reviewed.
"""
from __future__ import annotations

import csv
import json
from pathlib import Path

import pytest

from tools import shot_review_sheet as srs


def _clip(tmp_path):
    c = tmp_path / "clip"
    (c / "_labeling").mkdir(parents=True)
    (c / "classified.json").write_text(json.dumps({"shots": [
        {"shot_id": 0, "frame": 60, "t_sec": 1.0, "track_id": 7,
         "hitter_side": "near", "shot_type": "drive", "is_volley": False},
        {"shot_id": 1, "frame": 300, "t_sec": 5.0, "track_id": 9,
         "hitter_side": "far", "shot_type": "dink", "is_volley": True}]}), encoding="utf-8")
    (c / "track_roles.json").write_text(json.dumps(
        {"roles": {"user": {"track_ids": [7]}, "opp_a": {"track_ids": [9]}}}), encoding="utf-8")
    (c / "court.json").write_text(json.dumps({"video": {"fps": 60.0}}), encoding="utf-8")
    return c


def test_rows_are_numbered_as_the_video_numbers_them(tmp_path):
    """The '#' in the sheet and the '#N' burned into the annotated video must be the same
    number, or a correction lands on the wrong shot. Both number shots sorted by FRAME."""
    rows = srs.rows_for(_clip(tmp_path))
    assert [r["n"] for r in rows] == [1, 2]
    assert rows[0]["hitter"] == "user" and rows[0]["our_type"] == "drive"
    assert rows[1]["our_volley"] == "yes"


def test_build_refuses_to_clobber_a_filled_in_review(tmp_path, monkeypatch):
    """The operator's time is the scarce input. A rebuild on top of a completed review would
    destroy hours of it silently -- which nearly happened, caught only because Excel had the
    file locked."""
    import sys
    from openpyxl import load_workbook
    c = _clip(tmp_path)
    out = c / "_labeling" / srs.OUT_NAME
    srs.build(c, out)
    wb = load_workbook(out)
    ws = wb.active
    hdr = next(r for r in range(1, 30) if ws.cell(row=r, column=1).value == "#")
    col = {str(ws.cell(row=hdr, column=i).value or "").strip(): i
           for i in range(1, ws.max_column + 1)}
    ws.cell(row=hdr + 2, column=col["CORRECT_TYPE"], value="drop")
    wb.save(out)
    monkeypatch.setattr(sys, "argv", ["x", str(c)])
    with pytest.raises(SystemExit) as e:
        srs.main([str(c)])
    assert "Refusing to overwrite" in str(e.value)
    # --force gets through
    srs.main([str(c), "--force"])


def test_blank_correction_means_agree(tmp_path):
    """Silence is agreement. A reviewed shot we got right is still a label, and dropping it
    would quietly bias the truth set towards our own mistakes."""
    c = _clip(tmp_path)
    out = c / "_labeling" / srs.OUT_NAME
    srs.build(c, out)
    srs.score(c, out)
    with (c / "_labeling" / srs.CSV_NAME).open(encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    assert [r["true_type"] for r in rows] == ["drive", "dink"]


def test_a_correction_overrides_and_a_missed_shot_is_added(tmp_path):
    from openpyxl import load_workbook
    c = _clip(tmp_path)
    out = c / "_labeling" / srs.OUT_NAME
    srs.build(c, out)
    wb = load_workbook(out)
    ws = wb.active
    hdr = next(r for r in range(1, 30) if ws.cell(row=r, column=1).value == "#")
    # Find the columns by HEADER. The layout has gained an ALREADY KNOWN column and then
    # NOT_A_SHOT / RALLY_END; a test pinned to a letter breaks on every such change and, worse,
    # would pass while writing into the wrong column.
    col = {str(ws.cell(row=hdr, column=i).value or "").strip(): i
           for i in range(1, ws.max_column + 1)}
    ws.cell(row=hdr + 2, column=col["CORRECT_TYPE"], value="drop")
    blank = next(r for r in range(hdr, ws.max_row + 1)
                 if str(ws.cell(row=r, column=1).value or "").startswith("SHOTS WE MISSED"))
    ws.cell(row=blank + 1, column=col["time"], value="0:09.50")   # never detected
    ws.cell(row=blank + 1, column=col["CORRECT_TYPE"], value="lob")
    wb.save(out)
    srs.score(c, out)
    with (c / "_labeling" / srs.CSV_NAME).open(encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    assert rows[0]["true_type"] == "drop", "the correction did not override our type"
    missed = [r for r in rows if "MISSED" in r["notes"]]
    assert len(missed) == 1 and missed[0]["true_type"] == "lob"
    assert int(missed[0]["frame"]) == int(round(9.5 * 60)), "missed shot lost its timestamp"


def test_output_name_is_globbed_by_the_scorer():
    """score_shot_types reads every labels*.csv in _labeling; a name outside that pattern
    would make the whole review invisible."""
    assert srs.CSV_NAME.startswith("labels") and srs.CSV_NAME.endswith(".csv")


def test_clock_round_trips():
    for t in (0.0, 9.5, 63.82, 301.05):
        assert abs(srs.parse_clock(srs.clock(t)) - t) < 0.01


def test_not_a_shot_column_is_read(tmp_path):
    """The operator had to bury these in free text last time, and a regex bug then read none
    of them -- 16 false positives lost. A column cannot be silently mis-parsed."""
    from openpyxl import load_workbook
    c = _clip(tmp_path)
    out = c / "_labeling" / srs.OUT_NAME
    srs.build(c, out)
    wb = load_workbook(out)
    ws = wb.active
    hdr = next(r for r in range(1, 30) if ws.cell(row=r, column=1).value == "#")
    col = {str(ws.cell(row=hdr, column=i).value or "").strip(): i
           for i in range(1, ws.max_column + 1)}
    assert "NOT_A_SHOT" in col and "RALLY_END" in col
    ws.cell(row=hdr + 2, column=col["NOT_A_SHOT"], value="y")
    wb.save(out)
    srs.score(c, out)
    with (c / "_labeling" / srs.CSV_NAME).open(encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    assert rows[0]["true_type"] == "not a shot"


def test_every_control_lands_on_the_column_its_header_names(tmp_path):
    """The layout has gained a column three times (ALREADY KNOWN, NOT_A_SHOT/RALLY_END, the
    rally). Each time a hard-coded letter put a dropdown or a read on the wrong column --
    silently, because a dropdown on the wrong column still looks like a working sheet."""
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    c = _clip(tmp_path)
    out = c / "_labeling" / srs.OUT_NAME
    srs.build(c, out)
    ws = load_workbook(out).active
    hdr = next(r for r in range(1, 30) if ws.cell(row=r, column=1).value == "#")
    col = {str(ws.cell(row=hdr, column=i).value or "").strip(): i
           for i in range(1, ws.max_column + 1)}
    want = {"CORRECT_TYPE": "drive", "CORRECT_VOLLEY": "yes",
            "NOT_A_SHOT": "y", "RALLY_END": "y"}
    ranges = {}
    for dv in ws.data_validations.dataValidation:
        for rng in str(dv.sqref).split():
            ranges[rng.split(str(hdr + 2))[0].rstrip("0123456789:")] = dv.formula1
    for name, sample in want.items():
        letter = get_column_letter(col[name])
        assert letter in ranges, f"no dropdown on {name} (column {letter})"
        assert sample in ranges[letter], f"{name}'s dropdown does not offer {sample!r}"


def test_the_operators_own_rally_windows_are_shown_as_context(tmp_path):
    """Their objection to the first labelling tool: "you can't determine a shot at the point
    of contact without seeing it in context of where it is coming from and where it is
    going." The rally column is that context, and it is THEIR window, never our inference."""
    from openpyxl import load_workbook
    import tools.truth_store as ts
    c = _clip(tmp_path)
    # the store is keyed by SOURCE VIDEO, so the clip has to say which one it came from
    (c / "ball.meta.json").write_text(json.dumps({"video_path": "V.mp4"}), encoding="utf-8")
    doc = ts.empty("V.mp4")
    doc["rally_truth"] = [{"start_t_sec": 0.5, "end_t_sec": 2.0, "server": "user",
                           "n_shots": 4, "source": "t"}]
    ts.save(doc)
    try:
        srs.build(c, c / "_labeling" / srs.OUT_NAME)
        ws = load_workbook(c / "_labeling" / srs.OUT_NAME).active
        hdr = next(r for r in range(1, 30) if ws.cell(row=r, column=1).value == "#")
        col = {str(ws.cell(row=hdr, column=i).value or "").strip(): i
               for i in range(1, ws.max_column + 1)}
        got = [str(ws.cell(row=hdr + 1 + i, column=col["your rally"]).value or "")
               for i in (1, 2)]
        assert got == ["1", "between points"]
    finally:
        ts.store_path("V.mp4").unlink(missing_ok=True)


def test_a_corrected_hitter_replaces_the_one_we_guessed(tmp_path):
    """There was nowhere to say "you credited the wrong player", so the operator said it in
    the notes -- "shot was by opponent on far side", "dink by partner", "shot was by
    partner, not user" -- and the importer, reading only structured columns, kept our wrong
    hitter all four times. Server attribution is scored against exactly that field, so the
    answer key itself was wrong where we most needed it right."""
    from openpyxl import load_workbook
    c = _clip(tmp_path)
    out = c / "_labeling" / srs.OUT_NAME
    srs.build(c, out)
    wb = load_workbook(out)
    ws = wb.active
    hdr = next(r for r in range(1, 30) if ws.cell(row=r, column=1).value == "#")
    col = {str(ws.cell(row=hdr, column=i).value or "").strip(): i
           for i in range(1, ws.max_column + 1)}
    assert "CORRECT_HITTER" in col and "CORRECT_SIDE" in col

    row = hdr + 2
    ours_hitter = str(ws.cell(row=row, column=col["hitter"]).value or "")
    ws.cell(row=row, column=col["CORRECT_HITTER"], value="partner")
    ws.cell(row=row, column=col["CORRECT_SIDE"], value="far")
    wb.save(out)
    srs.score(c, out)

    with (c / "_labeling" / srs.CSV_NAME).open(encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    got = rows[0]
    assert got["hitter_role"] == "partner", (
        f"the correction was ignored; still {got['hitter_role']!r} (we said {ours_hitter!r})")
    assert got["hitter_side"] == "far"
    assert "corrected" in got["notes"]


def test_correcting_only_the_hitter_keeps_our_side(tmp_path):
    """The two are separate answers: the operator often knows WHO hit it without disputing
    which end of the court they were on. Overwriting side with a blank would erase a fact
    nobody questioned."""
    from openpyxl import load_workbook
    c = _clip(tmp_path)
    out = c / "_labeling" / srs.OUT_NAME
    srs.build(c, out)
    wb = load_workbook(out)
    ws = wb.active
    hdr = next(r for r in range(1, 30) if ws.cell(row=r, column=1).value == "#")
    col = {str(ws.cell(row=hdr, column=i).value or "").strip(): i
           for i in range(1, ws.max_column + 1)}
    row = hdr + 2
    ours_side = str(ws.cell(row=row, column=col["side"]).value or "")
    ws.cell(row=row, column=col["CORRECT_HITTER"], value="opp_a")
    wb.save(out)
    srs.score(c, out)
    with (c / "_labeling" / srs.CSV_NAME).open(encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    assert rows[0]["hitter_role"] == "opp_a"
    assert rows[0]["hitter_side"] == ours_side
